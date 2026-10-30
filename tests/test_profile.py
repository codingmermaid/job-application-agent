import io
import re

import pytest
from docx import Document
from fastapi.testclient import TestClient
from langgraph.checkpoint.memory import InMemorySaver

from jobagent.agents.profile_builder import ProfileDraft, extract_profile
from jobagent.documents import extract_text
from jobagent.llm import ScriptedLLM
from jobagent.web.app import content_disposition, create_app

PROFILE = ProfileDraft(
    name="Test Person", email="test@example.com", location="Belgrade, Serbia", country="Serbia",
    work_authorization="", portfolio_url="", headline="Analyst with SQL and Power BI experience.",
    facts=[f"Fact number {i}." for i in range(1, 10)] + ["  "],  # a blank line the agent left in
)


def docx_bytes(*paragraphs: str) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# reading CVs

def test_extract_text_from_word_and_text_files():
    assert "SQL" in extract_text("cv.docx", docx_bytes("Jane Doe", "Skills: SQL, Excel"))
    assert extract_text("cv.txt", "Hello CV".encode()) == "Hello CV"


def test_extract_text_rejects_bad_files():
    with pytest.raises(ValueError, match="Unsupported"):
        extract_text("cv.png", b"...")
    with pytest.raises(ValueError, match="2 MB"):
        extract_text("cv.txt", b"x" * (3 * 1024 * 1024))
    with pytest.raises(ValueError, match="No text"):
        extract_text("cv.txt", b"   ")


# profile builder agent

def test_extract_profile_numbers_facts_and_drops_blanks():
    candidate = extract_profile(ScriptedLLM({ProfileDraft: PROFILE}), "a CV")
    assert [f.id for f in candidate.facts] == [f"F{i}" for i in range(1, 10)]
    assert candidate.work_authorization == "Not stated"
    assert candidate.id.startswith("profile-")


# web: profile, runs and downloads

def test_download_names_with_accents_are_valid_headers():
    header = content_disposition("Mila-Jovanović-resume.docx")
    header.encode("latin-1")  # would raise if the header were invalid
    assert 'filename="Mila-Jovanovic-resume.docx"' in header
    assert "filename*=UTF-8''Mila-Jovanovi%C4%87-resume.docx" in header


def test_sample_candidate_downloads_work(world, engine, scripted_responses):
    client = TestClient(create_app(world=world, llm=ScriptedLLM(scripted_responses), engine=engine,
                                   checkpointer=InMemorySaver()))
    run_id = re.search(r"/runs/(\w+)", str(client.post("/runs", data={"request": "analyst"}).url)).group(1)
    for _ in range(20):
        if "Choose up to" in client.get(f"/runs/{run_id}/progress").text:
            break
    client.post(f"/runs/{run_id}/select", data={"job_ids": ["j01"]})
    for _ in range(20):
        if "Review and approve" in client.get(f"/runs/{run_id}/progress").text:
            break
    response = client.get(f"/runs/{run_id}/apps/j01/cover-letter.docx")
    assert response.status_code == 200 and "Jovanovi" in response.headers["content-disposition"]

def test_profile_flow_runs_and_downloads(world, engine, scripted_responses):
    scripted_responses[ProfileDraft] = PROFILE
    client = TestClient(create_app(world=world, llm=ScriptedLLM(scripted_responses), engine=engine,
                                   checkpointer=InMemorySaver()))
    assert "fictional sample" in client.get("/").text

    review = client.post("/profile/extract", data={"cv_text": "My CV text"}).text
    assert "Fact number 9." in review and 'value="Test Person"' in review

    facts = "\n".join(f"Fact number {i}." for i in range(1, 10))
    client.post("/profile/save", data={"name": "Test Person", "email": "test@example.com", "location": "Belgrade",
                                       "country": "Serbia", "facts": facts, "no_experience_with": "dbt, R"})
    home = client.get("/").text
    assert "Test Person" in home and "fictional sample" not in home

    run_id = re.search(r"/runs/(\w+)", str(client.post("/runs", data={"request": "analyst roles"}).url)).group(1)
    for _ in range(20):
        if "Choose up to" in client.get(f"/runs/{run_id}/progress").text:
            break
    client.post(f"/runs/{run_id}/select", data={"job_ids": ["j01"]})
    for _ in range(20):
        if "Review and approve" in client.get(f"/runs/{run_id}/progress").text:
            break

    # Changing the profile now must not change the run already in progress.
    client.post("/profile/reset")
    state = client.app.state.runner.view(run_id)["state"]
    assert state["candidate"]["name"] == "Test Person"

    resume = client.get(f"/runs/{run_id}/apps/j01/resume.docx")
    assert resume.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Test-Person-resume-Nimbus-Ledger.docx" in resume.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(resume.content)).paragraphs)
    assert "Test Person" in text and "Power BI" in text

    letter = client.get(f"/runs/{run_id}/apps/j01/cover-letter.txt").text
    assert "Nimbus Ledger" in letter and letter.rstrip().endswith("Test Person")
    assert client.get(f"/runs/{run_id}/apps/j99/resume.docx").status_code == 404
    assert client.get(f"/runs/{run_id}/apps/j01/resume.exe").status_code == 404
