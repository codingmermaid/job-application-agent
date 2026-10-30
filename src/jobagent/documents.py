"""Reading uploaded CVs and producing downloadable application documents."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .agents.submission import render_cover_letter, render_resume
from .models import Candidate

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_CV_CHARS = 20_000
SUPPORTED = (".pdf", ".docx", ".txt", ".md")


def extract_text(filename: str, data: bytes) -> str:
    """Plain text from an uploaded CV (PDF, Word, or text)."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("That file is larger than 2 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    elif suffix == ".docx":
        doc = Document(io.BytesIO(data))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        text = "\n".join([p.text for p in doc.paragraphs] + cells)
    elif suffix in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type. Use one of: {', '.join(SUPPORTED)}")
    text = text.strip()
    if not text:
        raise ValueError("No text found in that file. Scanned PDFs need to be pasted as text instead.")
    return text[:MAX_CV_CHARS]


def _docx(paragraphs: list[str], *, heading: str | None = None) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs:
        if p.startswith("- "):
            doc.add_paragraph(p[2:], style="List Bullet")
        else:
            doc.add_paragraph(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def resume_text(candidate: Candidate, draft: dict) -> str:
    return render_resume(candidate, draft)


def cover_letter_text(candidate: Candidate, draft: dict) -> str:
    return render_cover_letter(candidate, draft)


def resume_docx(candidate: Candidate, draft: dict) -> bytes:
    lines = render_resume(candidate, draft).splitlines()
    body = [line for line in lines[1:] if line.strip()]  # first line becomes the heading
    return _docx(body, heading=candidate.name)


def cover_letter_docx(candidate: Candidate, draft: dict) -> bytes:
    return _docx([p for p in render_cover_letter(candidate, draft).split("\n\n") if p.strip()])
